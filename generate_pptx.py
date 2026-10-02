from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 Widescreen
prs.slide_height = Inches(7.5)

# Color Palette (Matching your website design)
BG_COLOR = RGBColor(15, 23, 42)      # Dark Navy/Charcoal
TEXT_COLOR = RGBColor(255, 255, 255) # White
ACCENT_COLOR = RGBColor(132, 204, 22) # Lime Green
SUB_COLOR = RGBColor(148, 163, 184)   # Light Gray

def set_slide_background(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # Blank layout
    set_slide_background(slide)
    
    # Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
    tf = txBox.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = ACCENT_COLOR
    
    # Subtitle
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(11), Inches(1.5))
    tf2 = txBox2.text_frame
    tf2.text = subtitle
    p2 = tf2.paragraphs[0]
    p2.font.size = Pt(20)
    p2.font.color.rgb = TEXT_COLOR

def add_content_slide(title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    
    # Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1))
    tf = txBox.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = ACCENT_COLOR
    
    # Bullets
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(11), Inches(5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        p = tf2.add_paragraph()
        p.text = f"• {bullet}"
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(15)

def add_table_slide(title, headers, rows, footnote=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    
    # Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11), Inches(1))
    tf = txBox.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = ACCENT_COLOR
    
    # Table
    rows_count = len(rows) + 1
    cols_count = len(headers)
    left = Inches(1)
    top = Inches(2)
    width = Inches(11)
    height = Inches(4)
    
    table = slide.shapes.add_table(rows_count, cols_count, left, top, width, height).table
    
    # Set Headers
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(18)
        p.font.color.rgb = ACCENT_COLOR
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 41, 59) # Slightly lighter dark
        
    # Set Rows
    for r, row_data in enumerate(rows):
        for c, cell_data in enumerate(row_data):
            cell = table.cell(r+1, c)
            cell.text = cell_data
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(16)
            p.font.color.rgb = TEXT_COLOR
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG_COLOR
            
    # Footnote
    if footnote:
        txBox2 = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11), Inches(0.5))
        tf2 = txBox2.text_frame
        tf2.text = footnote
        p = tf2.paragraphs[0]
        p.font.size = Pt(12)
        p.font.color.rgb = SUB_COLOR

# --- BUILDING THE SLIDES ---

# Slide 1: Title
add_title_slide(
    "2026 India & International Market Pricing Study",
    "A practical benchmark of 2026 web-development market trends and the strategic positioning of PSDigiLabs."
)

# Slide 2: Market Landscape
add_content_slide(
    "The 2026 Market Landscape",
    [
        "India: Market has matured. Professional custom builds command a premium (₹2L+ for custom apps).",
        "International: Custom web apps in US start at $25,000+, UK at £15,000+.",
        "Strategic Implication: A massive gap exists between 'cheap template' and 'premium enterprise'.",
        "PSDigiLabs occupies the high-value middle ground: modern engineering, lean delivery, transparent scope."
    ]
)

# Slide 3: India Table
add_table_slide(
    "Indicative India Market Ranges (2026)",
    ["Service Category", "Revised India Range (2026)"],
    [
        ["Landing Page", "₹8,000 – ₹25,000"],
        ["Basic Business Website", "₹20,000 – ₹60,000"],
        ["Professional Business Website", "₹50,000 – ₹1,50,000"],
        ["E-commerce Website", "₹60,000 – ₹4,00,000+"],
        ["Custom Web Application", "₹2,00,000 – ₹15,00,000+"]
    ],
    "Footnote: Excludes PSDigiLabs internal pricing. Ranges are indicative planning benchmarks."
)

# Slide 4: US Table
add_table_slide(
    "Indicative United States Market Ranges (2026)",
    ["Project Type", "Revised US Range (2026)"],
    [
        ["Small Business Website", "$3,000 – $10,000"],
        ["E-commerce Website", "$5,000 – $25,000+"],
        ["Custom Web App", "$25,000 – $150,000+"],
        ["Website Maintenance", "$50 – $500 / month"]
    ]
)

# Slide 5: UK Table
add_table_slide(
    "Indicative United Kingdom Market Ranges (2026)",
    ["Project Type", "Revised UK Range (2026)"],
    [
        ["Business Website", "£500 – £3,000"],
        ["E-commerce Website", "£3,000 – £25,000+"],
        ["Custom Web App", "£15,000 – £100,000+"],
        ["Website Maintenance", "£50 – £300 / month"]
    ]
)

# Slide 6: Canada Table
add_table_slide(
    "Indicative Canada Market Ranges (2026)",
    ["Project Type", "Revised Canada Range (2026)"],
    [
        ["Business Website", "CA$2,500 – CA$10,000"],
        ["E-commerce Website", "CA$5,000 – CA$25,000+"],
        ["Custom Web App", "CA$15,000 – CA$40,000+"],
        ["Website Maintenance", "CA$100 – CA$500 / month"]
    ]
)

# Slide 7: Australia Table
add_table_slide(
    "Indicative Australia Market Ranges (2026)",
    ["Project Type", "Revised Australia Range (2026)"],
    [
        ["Business Website", "A$3,000 – A$15,000"],
        ["E-commerce Website", "A$8,000 – A$25,000+"],
        ["Custom Web App", "A$20,000 – A$300,000+"],
        ["Website Maintenance", "A$100 – A$500 / month"]
    ]
)

# Slide 8: Value Proposition
add_content_slide(
    "Why International Clients Choose an India-Based Engineering Partner",
    [
        "Lean Delivery: Direct communication, lower agency overhead, faster decisions.",
        "Modern Stack: Next.js, React, TypeScript, PostgreSQL, Tailwind.",
        "Quality Built-In: Testing and QA integrated from day one.",
        "Scope Discipline: Clear boundaries. 'Starting price' is an entry point, not an open promise."
    ]
)

# Slide 9: Customer Message
add_content_slide(
    "Sell the Outcome and the Engineering Model — Not a Discount",
    [
        "\"International-quality digital product delivery with the cost efficiency of an India-based engineering partner.\"",
        "Avoid: 'Cheap Indian development' or percentage-discount claims.",
        "Emphasize: Clear scope, modern engineering, rigorous QA, seamless deployment.",
        "Prove: Live work, detailed case studies, documented processes, client testimonials."
    ]
)

# Slide 10: Pricing Path
add_content_slide(
    "Use Entry-Level Pricing as a Strategy, Then Raise It with Evidence",
    [
        "Phase 1: Portfolio Building (Current) – Competitive rates to attract early clients.",
        "Phase 2: Value Expansion (6–12 Months) – Increase bands as case studies prove outcomes.",
        "Phase 3: Premium Partnership (12+ Months) – Established rates based on proven delivery.",
        "Note: Maintenance and support should be priced separately as a recurring retainer."
    ]
)

# Slide 11: Scope Protection
add_content_slide(
    "Competitive Starting Prices Require Disciplined Scoping",
    [
        "Define the Base Scope: List pages, roles, integrations, content, design rounds.",
        "Price Add-ons Separately: Extra integrations, data migration, copywriting, advanced SEO.",
        "Limit Revision Rounds: Prevent open-ended design and content cycles.",
        "Use Milestones: Deposit → Design/Architecture → Build → UAT → Launch.",
        "Separate Support: Clarify warranty/fix period versus ongoing maintenance.",
        "Reprice Complexity: A 'starting from' price is not a fixed quote for every project."
    ]
)

# Slide 12: Methodology
add_content_slide(
    "Methodology & Sources",
    [
        "Method: Sourced from publicly available 2026 pricing guides, agency websites, and freelance platforms (India, US, UK, Canada, Australia).",
        "Disclaimer: Categories are normalized at a high level. Actual scope, design depth, and integrations can change quotes materially.",
        "This is a comparative strategy study, not a promise that two projects with the same label have equivalent scope."
    ]
)

# Save the file
prs.save('PSDigiLabs_2026_Market_Study.pptx')
print("Presentation created successfully: PSDigiLabs_2026_Market_Study.pptx")